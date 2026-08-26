"""Shared row-iteration for bulk teacher/student upload (.csv or .xlsx) -
column-name-based (via the header row) so column order in the uploaded
file doesn't matter. Parsing only; the caller owns validation, seat-limit
checks, and Firebase provisioning per row.
"""

import csv
import io
from dataclasses import dataclass
from typing import Literal

from openpyxl import load_workbook

from shared.errors import ValidationError

MAX_ROWS = 500
MAX_FILE_BYTES = 2 * 1024 * 1024


@dataclass
class BulkRowResult:
    row: int
    status: Literal["created", "error"]
    email: str | None = None
    temp_password: str | None = None
    error: str | None = None


def parse_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    """Returns one dict per data row, keyed by lowercased header cell.
    Raises ValidationError (surfaced as a 422, per the standard error
    envelope) for anything that isn't a well-formed .csv/.xlsx, is empty,
    or exceeds the row/size caps - these are rejected before any Firebase
    account is created, not discovered mid-batch.
    """
    if len(content) > MAX_FILE_BYTES:
        raise ValidationError("File is too large - the limit is 2MB.")

    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        rows = _parse_csv(content)
    elif lower_name.endswith(".xlsx"):
        rows = _parse_xlsx(content)
    else:
        raise ValidationError("Only .csv or .xlsx files are supported.")

    if not rows:
        raise ValidationError("The file has no data rows.")
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"The file has {len(rows)} rows - the limit is {MAX_ROWS} per upload.")
    return rows


def _parse_csv(content: bytes) -> list[dict[str, str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError("The CSV file isn't valid UTF-8 text.") from exc
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise ValidationError("The CSV file has no header row.")
    return [
        {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        for row in reader
    ]


def _parse_xlsx(content: bytes) -> list[dict[str, str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("Couldn't read the .xlsx file - is it corrupted?") from exc

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell or "").strip().lower() for cell in next(rows_iter)]
    except StopIteration:
        raise ValidationError("The spreadsheet has no header row.") from None

    rows = []
    for raw_row in rows_iter:
        if all(cell is None or str(cell).strip() == "" for cell in raw_row):
            continue  # skip fully blank rows (trailing blank rows are common in exported sheets)
        rows.append({header[i]: str(cell).strip() if cell is not None else "" for i, cell in enumerate(raw_row) if i < len(header)})
    return rows
